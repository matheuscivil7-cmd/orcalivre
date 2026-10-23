"""Robô do OrçaLivre: baixa o SINAPI direto da Caixa e gera os arquivos que o app consulta.

Uso:
  python robo/atualizar_sinapi.py                 # baixa o mês mais recente, se ainda não existir
  python robo/atualizar_sinapi.py --meses 3       # garante os 3 meses mais recentes
  python robo/atualizar_sinapi.py --arquivo X.zip # processa um pacote já baixado (zip ou xlsx)
  python robo/atualizar_sinapi.py --forcar        # refaz mesmo que o mês já exista

Saída em dados/sinapi/AAAA-MM/:
  cat.json  — insumos [código, descrição, unidade, tipo], composições
              [código, descrição, unidade, grupo, caderno, itens(t,código,coef...)], grupos, cadernos
  UF.json   — {SD|CD|SE: {i: preços dos insumos, c: custos das composições, a: índices com preço de SP}}
e atualiza dados/bases.json. Se o layout da planilha mudar ou os números não fecharem,
o robô para com erro e não grava nada.
"""
import argparse, io, json, os, re, shutil, sys, tempfile, unicodedata, zipfile
import urllib.request, http.cookiejar

import openpyxl

UFS = ["AC","AL","AM","AP","BA","CE","DF","ES","GO","MA","MG","MS","MT","PA","PB","PE","PI","PR","RJ","RN","RO","RR","RS","SC","SE","SP","TO"]
REGIMES = {"SD": ("ISD", "CSD"), "CD": ("ICD", "CCD"), "SE": ("ISE", "CSE")}
RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DADOS = os.environ.get("ORCALIVRE_DADOS", os.path.join(RAIZ, "dados"))
MANTER_MESES = 12
CAIXA = "https://www.caixa.gov.br"
CATEGORIA = 888  # "SINAPI - Relatórios mensais - a partir de 2025"


def log(*a):
    print(*a, flush=True)


def norm(s):
    return unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().upper().strip()


def txt(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "")).strip()


def numero(v):
    if v is None or v == "" or v == "-":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("R$", "").replace(" ", "")
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def codigo(v):
    """Códigos do SINAPI são inteiros; nas composições a célula pode ser =HYPERLINK(...;104658)."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        n = int(v)
        return str(n) if n > 0 and n == v else None
    s = str(v).strip()
    if re.fullmatch(r"\d+(\.0+)?", s):
        n = int(float(s))
        return str(n) if n > 0 else None
    m = re.search(r"[;,]\s*\"?(\d+)\"?\s*\)\s*$", s)
    if m and s.upper().lstrip("=").startswith("HYPERLINK"):
        return m.group(1)
    return None


def tipo_insumo(cl, u):
    x = norm(cl)
    if "MAO DE OBRA" in x:
        return "MO"
    if "EQUIP" in x:
        return "EQ"
    if "MATERIAL" in x:
        return "MAT"
    if "SERVI" in x:
        return "SV"
    if x:
        return "OUT"
    return "EQ" if norm(u) in ("CHP", "CHI") else "MAT"


def exigir(cond, msg):
    if not cond:
        raise SystemExit(f"ERRO: layout do SINAPI inesperado — {msg}. Nada foi gravado.")


# ---------------------------------------------------------------- download
class Caixa:
    def __init__(self):
        self.op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
        self.op.addheaders = [("User-Agent", "Mozilla/5.0 (OrcaLivre robo SINAPI)")]
        self.op.open(f"{CAIXA}/poder-publico/modernizacao-gestao/sinapi/Paginas/default.aspx", timeout=60).read()

    def pacotes(self):
        url = (f"{CAIXA}/_api/web/lists/getbytitle('Downloads')/Items?$select=Title,FileLeafRef,EncodedAbsUrl,Modified,"
               f"Categoria/ID&$expand=Categoria&$filter=Categoria/ID%20eq%20{CATEGORIA}&$top=100&$orderby=Modified%20desc")
        rq = urllib.request.Request(url, headers={"Accept": "application/json;odata=nometadata"})
        dados = json.loads(self.op.open(rq, timeout=60).read().decode("utf-8"))
        achados = {}
        for it in dados.get("value", []):
            m = re.search(r"SINAPI-(\d{4})-(\d{2})-formato-xlsx(_Retificacao\d*)?\.zip$", it.get("FileLeafRef", ""), re.I)
            if not m:
                continue
            comp = f"{m.group(1)}-{m.group(2)}"
            # retificação mais recente prevalece
            if comp not in achados or it.get("Modified", "") > achados[comp]["Modified"]:
                achados[comp] = {"url": it["EncodedAbsUrl"].replace("http:", "https:"), "Modified": it.get("Modified", ""), "nome": it["FileLeafRef"]}
        return dict(sorted(achados.items(), reverse=True))

    def baixar(self, url, destino):
        with self.op.open(url, timeout=600) as r, open(destino, "wb") as f:
            shutil.copyfileobj(r, f)
        exigir(os.path.getsize(destino) > 1_000_000, f"pacote pequeno demais ({os.path.getsize(destino)} bytes)")


def planilha_referencia(caminho, pasta):
    if caminho.lower().endswith(".xlsx"):
        return caminho
    with zipfile.ZipFile(caminho) as z:
        nomes = [n for n in z.namelist() if n.lower().endswith(".xlsx")]
        alvo = [n for n in nomes if "REFER" in norm(n.encode("cp437", "ignore").decode("utf-8", "ignore")) or "REFER" in norm(n)]
        exigir(alvo or nomes, "o .zip não contém planilha .xlsx")
        nome = (alvo or nomes)[0]
        destino = os.path.join(pasta, "referencia.xlsx")
        with z.open(nome) as src, open(destino, "wb") as dst:
            shutil.copyfileobj(src, dst)
        return destino


# ---------------------------------------------------------------- leitura
def linhas(wb, aba):
    exigir(aba in wb.sheetnames, f"aba {aba} não encontrada")
    return [list(r) for r in wb[aba].iter_rows(values_only=True)]


def competencia(rows):
    for r in rows[:8]:
        if r and norm(r[0]).startswith("MES DE REFER"):
            m = re.match(r"(\d{2})/(\d{4})", txt(r[1]))
            exigir(m, f"competência inválida '{r[1]}'")
            return f"{m.group(2)}-{m.group(1)}"
    exigir(False, "cabeçalho 'Mês de Referência' ausente")


def emissao(rows):
    for r in rows[:8]:
        if r and norm(r[0]).startswith("DATA DE EMISS"):
            return txt(r[1])
    return txt(rows[3][1]) if len(rows) > 3 and rows[3] and len(rows[3]) > 1 else ""


def achar_cabecalho(rows, *padroes):
    for i, r in enumerate(rows[:20]):
        cab = [norm(c) for c in (r or [])]
        if all(any(re.search(p, c) for c in cab) for p in padroes):
            return i, cab
    exigir(False, f"cabeçalho com {padroes} não encontrado")


def colunas_uf(rows, h):
    """Posição de cada UF: procura a linha (cabeçalho ou acima) que contém as 27 siglas."""
    for i in range(max(0, h - 3), h + 1):
        r = [norm(c) for c in (rows[i] or [])]
        pos = {uf: r.index(uf) for uf in UFS if uf in r}
        if len(pos) == len(UFS):
            return pos
    exigir(False, "colunas das 27 UFs não encontradas")


def ler_insumos(wb, aba):
    rows = linhas(wb, aba)
    h, cab = achar_cabecalho(rows, r"CODIGO DO INSUMO|^CODIGO$", r"DESCRI", r"UNIDADE")
    ci = next(j for j, c in enumerate(cab) if re.search(r"CODIGO", c))
    cd = next(j for j, c in enumerate(cab) if "DESCRI" in c)
    cu = next(j for j, c in enumerate(cab) if "UNIDADE" in c)
    ccl = next((j for j, c in enumerate(cab) if "CLASSIF" in c), 0)
    pos = colunas_uf(rows, h)
    out = {}
    for r in rows[h + 1:]:
        if not r or len(r) <= ci:
            continue
        c = codigo(r[ci])
        if not c:
            continue
        out[c] = {"d": txt(r[cd]), "u": txt(r[cu]), "cl": txt(r[ccl]), "p": {uf: numero(r[j]) if j < len(r) else None for uf, j in pos.items()}}
    return rows, out


def ler_composicoes(wb, wbf, aba):
    rows = linhas(wb, aba)
    h, cab = achar_cabecalho(rows, r"CODIGO", r"DESCRI", r"UNIDADE")
    ci = next(j for j, c in enumerate(cab) if "CODIGO" in c)
    cd = next(j for j, c in enumerate(cab) if "DESCRI" in c)
    cu = next(j for j, c in enumerate(cab) if "UNIDADE" in c)
    pos = colunas_uf(rows, h)
    formulas = None
    out = {}
    for k, r in enumerate(rows[h + 1:]):
        if not r or len(r) <= ci:
            continue
        c = codigo(r[ci])
        if not c:
            if formulas is None:  # o código costuma vir num HYPERLINK cujo valor exibido é 0
                formulas = [row[ci] if len(row) > ci else None for row in wbf[aba].iter_rows(values_only=True)]
            c = codigo(formulas[h + 1 + k]) if h + 1 + k < len(formulas) else None
        if not c:
            continue
        out[c] = {"d": txt(r[cd]), "u": txt(r[cu]), "g": txt(r[0]), "p": {uf: numero(r[j]) if j < len(r) else None for uf, j in pos.items()}}
    return out


def ler_analitico(wb):
    aba = next((n for n in wb.sheetnames if norm(n).startswith("ANALITICO") and "CUSTO" not in norm(n)), None)
    exigir(aba, "aba Analítico não encontrada")
    rows = linhas(wb, aba)
    h, cab = achar_cabecalho(rows, r"CODIGO DA COMPOSI", r"TIPO ITEM", r"COEFICIENTE")
    j = {k: next(i for i, c in enumerate(cab) if re.search(p, c)) for k, p in
         {"comp": r"CODIGO DA COMPOSI", "tipo": r"TIPO ITEM", "item": r"CODIGO DO ITEM", "desc": r"DESCRI", "und": r"UNIDADE", "coef": r"COEFICIENTE"}.items()}
    est, extra = {}, {}
    for r in rows[h + 1:]:
        if not r:
            continue
        comp = codigo(r[j["comp"]])
        if not comp:
            continue
        tp = norm(r[j["tipo"]])
        if not tp:
            extra.setdefault(("C", comp), (txt(r[j["desc"]]), txt(r[j["und"]]), txt(r[0])))
            continue
        item, q = codigo(r[j["item"]]), numero(r[j["coef"]])
        if not item or q is None:
            continue
        t = "C" if tp.startswith("COMP") else "I"
        est.setdefault(comp, []).append((t, item, q))
        extra.setdefault((t, item), (txt(r[j["desc"]]), txt(r[j["und"]]), ""))
    return est, extra


def limpo(v):
    return None if v in (None, 0, 0.0) else round(float(v), 2)


# ---------------------------------------------------------------- geração
def cadernos_anteriores(mes):
    """Reaproveita os links dos cadernos técnicos do mês anterior mais próximo (a planilha não os traz)."""
    pasta = os.path.join(DADOS, "sinapi")
    if not os.path.isdir(pasta):
        return {}
    for m in sorted((x for x in os.listdir(pasta) if re.fullmatch(r"\d{4}-\d{2}", x)), reverse=True):
        if m == mes and not os.path.exists(os.path.join(pasta, m, "cat.json")):
            continue
        try:
            cat = json.load(open(os.path.join(pasta, m, "cat.json"), encoding="utf-8"))
            return {r[0]: cat["ct"][r[4]] for r in cat["comp"] if cat["ct"][r[4]]}
        except Exception:
            continue
    return {}


def processar(xlsx, forcar=False):
    log("Abrindo planilha…")
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    wbf = openpyxl.load_workbook(xlsx, read_only=True, data_only=False)
    rows_isd, ins = ler_insumos(wb, "ISD")
    mes = competencia(rows_isd)
    emis = emissao(rows_isd)
    pasta = os.path.join(DADOS, "sinapi", mes)
    if os.path.exists(os.path.join(pasta, "cat.json")) and not forcar:
        log(f"{mes} já existe; nada a fazer.")
        return None
    log(f"Competência {mes} (emissão {emis}); {len(ins)} insumos")
    precos_i = {"SD": ins, "CD": ler_insumos(wb, "ICD")[1], "SE": ler_insumos(wb, "ISE")[1]}
    precos_c = {r: ler_composicoes(wb, wbf, aba) for r, (_, aba) in REGIMES.items()}
    comp = precos_c["SD"]
    log(f"{len(comp)} composições")
    est, extra = ler_analitico(wb)
    log(f"Analítico: {sum(len(v) for v in est.values())} itens em {len(est)} composições")
    wb.close(); wbf.close()

    # itens que só aparecem no analítico (sem preço / em estudo) entram no catálogo sem preço
    for (t, c), (d, u, g) in extra.items():
        if t == "I" and c not in ins:
            ins[c] = {"d": d, "u": u, "cl": "", "p": {}}
        if t == "C" and c not in comp:
            comp[c] = {"d": d, "u": u, "g": g, "p": {}}

    lst_i = sorted(ins.items(), key=lambda kv: kv[1]["d"])
    lst_c = sorted(comp.items(), key=lambda kv: (kv[1]["g"], kv[1]["d"]))
    grupos, cadernos = [], [""]
    gi, cti = {}, {"": 0}
    ct_ant = cadernos_anteriores(mes)
    def ctidx(c):
        u = ct_ant.get(c, "")
        if u not in cti:
            cti[u] = len(cadernos); cadernos.append(u)
        return cti[u]
    def gidx(g):
        if g not in gi:
            gi[g] = len(grupos); grupos.append(g)
        return gi[g]
    cat = {"m": mes,
           "ins": [[c, v["d"], v["u"], tipo_insumo(v["cl"], v["u"])] for c, v in lst_i],
           "comp": [[c, v["d"], v["u"], gidx(v["g"]), ctidx(c), [x for it in est.get(c, []) for x in it]] for c, v in lst_c],
           "g": grupos, "ct": cadernos}

    # validação: com os preços de SP, a soma do analítico precisa reproduzir os custos oficiais
    pi = {c: (precos_i["SD"].get(c, {}).get("p", {}) or {}).get("SP") for c, _ in lst_i}
    pc = {c: (precos_c["SD"].get(c, {}).get("p", {}) or {}).get("SP") for c, _ in lst_c}
    ok = tot = 0
    for c, _ in lst_c:
        oficial = pc.get(c)
        if not oficial or c not in est:
            continue
        soma = sum(((pi if t == "I" else pc).get(i) or 0) * q for t, i, q in est[c])
        tot += 1
        ok += abs(soma - oficial) / oficial < 0.01
    taxa = ok / tot if tot else 0
    log(f"Conferência SP: {ok}/{tot} composições reproduzidas ({taxa:.1%})")
    exigir(len(lst_i) > 4000 and len(lst_c) > 8000, "quantidade de insumos/composições muito baixa")
    exigir(taxa >= 0.9, f"só {taxa:.1%} das composições de SP fecham com o analítico")

    tmp = pasta + ".tmp"
    shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
    with open(os.path.join(tmp, "cat.json"), "w", encoding="utf-8") as f:
        json.dump(cat, f, ensure_ascii=False, separators=(",", ":"))
    for uf in UFS:
        o = {}
        for r in REGIMES:
            vi = [limpo((precos_i[r].get(c, {}).get("p") or {}).get(uf)) for c, _ in lst_i]
            sp = [limpo((precos_i[r].get(c, {}).get("p") or {}).get("SP")) for c, _ in lst_i]
            a = []
            if uf != "SP":
                for k, v in enumerate(vi):
                    if v is None and sp[k] is not None:
                        vi[k] = sp[k]; a.append(k)
            vc = [limpo((precos_c[r].get(c, {}).get("p") or {}).get(uf)) for c, _ in lst_c]
            o[r] = {"i": vi, "c": vc, "a": a}
        with open(os.path.join(tmp, f"{uf}.json"), "w", encoding="utf-8") as f:
            json.dump(o, f, separators=(",", ":"))
    shutil.rmtree(pasta, ignore_errors=True); os.replace(tmp, pasta)
    return {"mes": mes, "ins": len(lst_i), "comp": len(lst_c), "an": sum(len(v) for v in est.values()), "emissao": emis, "conferencia_sp": round(taxa, 4)}


def atualizar_manifesto(novos):
    caminho = os.path.join(DADOS, "bases.json")
    man = json.load(open(caminho, encoding="utf-8")) if os.path.exists(caminho) else {}
    s = man.setdefault("sinapi", {"ufs": UFS, "info": {}, "fonte": "SINAPI – Caixa Econômica Federal (relatórios mensais XLSX)"})
    s["ufs"] = UFS
    for n in novos:
        s["info"][n["mes"]] = {k: v for k, v in n.items() if k != "mes"}
    pasta = os.path.join(DADOS, "sinapi")
    meses = sorted([m for m in os.listdir(pasta) if re.fullmatch(r"\d{4}-\d{2}", m) and os.path.exists(os.path.join(pasta, m, "cat.json"))], reverse=True)
    for velho in meses[MANTER_MESES:]:
        shutil.rmtree(os.path.join(pasta, velho), ignore_errors=True)
        s["info"].pop(velho, None)
    s["meses"] = meses[:MANTER_MESES]
    import datetime
    man["gerado"] = datetime.date.today().isoformat()
    with open(caminho, "w", encoding="utf-8") as f:
        json.dump(man, f, ensure_ascii=False, indent=1)
    log("Meses disponíveis:", ", ".join(s["meses"]))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--meses", type=int, default=1)
    ap.add_argument("--arquivo")
    ap.add_argument("--forcar", action="store_true")
    a = ap.parse_args()
    novos = []
    with tempfile.TemporaryDirectory() as tmp:
        if a.arquivo:
            r = processar(planilha_referencia(a.arquivo, tmp), a.forcar)
            if r: novos.append(r)
        else:
            cx = Caixa()
            pacotes = cx.pacotes()
            exigir(pacotes, "a lista de downloads da Caixa não trouxe pacotes XLSX do SINAPI")
            log("Pacotes na Caixa:", ", ".join(pacotes))
            for mes, p in list(pacotes.items())[:max(1, a.meses)]:
                if os.path.exists(os.path.join(DADOS, "sinapi", mes, "cat.json")) and not a.forcar:
                    log(f"{mes} já está no site."); continue
                log(f"Baixando {p['nome']}…")
                z = os.path.join(tmp, f"{mes}.zip"); cx.baixar(p["url"], z)
                r = processar(planilha_referencia(z, tmp), a.forcar)
                os.remove(z)
                if r: novos.append(r)
    if novos:
        atualizar_manifesto(novos)
    print("NOVOS=" + ",".join(n["mes"] for n in novos))
    gh = os.environ.get("GITHUB_OUTPUT")
    if gh:
        with open(gh, "a") as f:
            f.write("novos=" + ",".join(n["mes"] for n in novos) + "\n")


# --- ajustes de leitura da planilha oficial ---
# A planilha da Caixa declara dimensões incorretas nas abas; no modo de leitura rápida isso
# faz o openpyxl enxergar só a primeira célula. reset_dimensions() força a leitura completa.
_load_orig = openpyxl.load_workbook


def _load_reset(*a, **k):
    wb = _load_orig(*a, **k)
    for ws in wb.worksheets:
        if hasattr(ws, "reset_dimensions"):
            ws.reset_dimensions()
    return wb


openpyxl.load_workbook = _load_reset

_pr_orig = planilha_referencia


def planilha_referencia(caminho, pasta):
    if caminho.lower().endswith(".zip"):
        with zipfile.ZipFile(caminho) as z:
            log("Arquivos no pacote:", z.namelist())
    return _pr_orig(caminho, pasta)


_cab_orig = achar_cabecalho


def achar_cabecalho(rows, *padroes):
    try:
        return _cab_orig(rows, *padroes)
    except SystemExit:
        log("Primeiras linhas da aba (diagnóstico):")
        for r in rows[:14]:
            log("  ", [str(c)[:30] for c in (r or [])[:8]])
        raise


if __name__ == "__main__":
    main()
